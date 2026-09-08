"""Core import and analysis helpers for the Guagua XHS data pipeline.

Raw creator exports stay in ``local-data/`` by default. Only normalized content
metadata and derived JSON are intended for the public GitHub Pages repository.
"""

from __future__ import annotations

import csv
import json
import math
import os
import re
import statistics
import tempfile
from collections import defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Any, Dict, Iterable, List, Mapping, Optional, Sequence, Tuple
from zoneinfo import ZoneInfo


SHANGHAI = ZoneInfo("Asia/Shanghai")
SCHEMA_VERSION = "1.0"
METRICS = (
    "impressions",
    "views",
    "likes",
    "favorites",
    "comments",
    "shares",
    "followers_gained",
)
ACCOUNT_METRICS = (
    "followers",
    "followers_delta",
    "impressions",
    "views",
    "profile_views",
    "likes",
    "favorites",
    "comments",
    "shares",
    "engagements",
    "new_followers",
)
WINDOWS = {"24h": 1, "72h": 3, "7d": 7, "30d": 30}


class PipelineError(RuntimeError):
    """User-facing pipeline failure."""


class DuplicateSnapshotError(PipelineError):
    """Raised when a snapshot date already exists with different content."""


def now_iso() -> str:
    return datetime.now(SHANGHAI).replace(microsecond=0).isoformat()


def today_iso() -> str:
    return datetime.now(SHANGHAI).date().isoformat()


def parse_date(value: Any, field: str = "date") -> Optional[str]:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y年%m月%d日", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            pass
    raise PipelineError(f"{field} 无法识别日期：{text!r}；请使用 YYYY-MM-DD。")


def to_number(value: Any) -> Optional[float]:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return float(value)
    if isinstance(value, (int, float)):
        if isinstance(value, float) and math.isnan(value):
            return None
        return float(value)
    text = str(value).strip().replace(",", "").replace("，", "")
    if not text or text.lower() in {"null", "none", "nan", "-", "—", "暂无"}:
        return None
    multiplier = 1.0
    if text.endswith("万"):
        multiplier, text = 10000.0, text[:-1]
    if text.endswith("+"):
        text = text[:-1]
    try:
        result = float(text) * multiplier
    except ValueError as exc:
        raise PipelineError(f"无法识别数值：{value!r}") from exc
    return int(result) if result.is_integer() else result


def _header_key(value: Any) -> str:
    text = str(value or "").strip().lower()
    return re.sub(r"[\s_\-（）()\/]+", "", text)


ALIASES = {
    "snapshot_date": ["snapshotdate", "快照日期", "采集日期", "统计日期", "日期", "date"],
    "content_id": ["contentid", "内容id", "内容编号", "作品编号"],
    "note_id": ["noteid", "笔记id", "笔记编号", "作品id"],
    "title": ["title", "标题", "笔记标题", "作品标题", "内容标题"],
    "publish_date": ["publishdate", "发布日期", "发布时间", "首发时间"],
    "followers": ["followers", "粉丝数", "总粉丝", "粉丝总数"],
    "followers_delta": ["followersdelta", "粉丝变化", "净涨粉", "粉丝增量"],
    "impressions": ["impressions", "曝光", "曝光数", "曝光量", "展示次数"],
    "views": ["views", "阅读", "阅读数", "阅读量", "观看", "观看次数", "播放量"],
    "profile_views": ["profileviews", "主页访问", "主页访问量", "主页浏览"],
    "likes": ["likes", "点赞", "点赞数", "点赞量"],
    "favorites": ["favorites", "收藏", "收藏数", "收藏量", "saves"],
    "comments": ["comments", "评论", "评论数", "评论量"],
    "shares": ["shares", "分享", "分享数", "分享量"],
    "engagements": ["engagements", "互动", "互动数", "互动量"],
    "followers_gained": ["followersgained", "笔记涨粉", "涨粉", "涨粉人数"],
    "new_followers": ["newfollowers", "新增粉丝", "新增关注", "新增涨粉"],
    "period_start": ["periodstart", "周期开始", "统计开始", "开始日期"],
    "period_end": ["periodend", "周期结束", "统计结束", "结束日期"],
    "metric_window": ["metricwindow", "aggregationtype", "指标口径", "统计口径", "数据口径"],
}
ALIAS_LOOKUP = {
    _header_key(alias): canonical for canonical, aliases in ALIASES.items() for alias in aliases
}


def canonicalize_row(row: Mapping[str, Any]) -> Dict[str, Any]:
    result: Dict[str, Any] = {}
    for raw_key, value in row.items():
        canonical = ALIAS_LOOKUP.get(_header_key(raw_key), str(raw_key).strip())
        if canonical and canonical not in result:
            result[canonical] = value
    return result


def _find_header_index(rows: Sequence[Sequence[Any]]) -> Optional[int]:
    scored = []
    for index, values in enumerate(rows[:25]):
        score = sum(_header_key(value) in ALIAS_LOOKUP for value in values if value not in (None, ""))
        scored.append((score, index))
    if not scored:
        return None
    score, index = max(scored)
    return index if score >= 2 else None


def read_csv_rows(path: Path) -> List[Dict[str, Any]]:
    last_error: Optional[Exception] = None
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            with path.open("r", encoding=encoding, newline="") as handle:
                return [canonicalize_row(row) for row in csv.DictReader(handle)]
        except UnicodeDecodeError as exc:
            last_error = exc
    raise PipelineError(f"无法读取 CSV 编码：{path}") from last_error


def _xhs_account_daily_rows(sheet_rows: Mapping[str, Sequence[Sequence[Any]]]) -> List[Dict[str, Any]]:
    """Convert XHS trend worksheets into normalized daily account rows."""
    rows_by_date: Dict[str, Dict[str, Any]] = {}
    for sheet_name, metric in (("曝光趋势", "impressions"), ("观看趋势", "views")):
        rows = sheet_rows.get(sheet_name) or []
        for values in rows[1:]:
            if len(values) < 2 or values[0] in (None, ""):
                continue
            snapshot_date = parse_date(values[0], f"{sheet_name} 日期")
            if not snapshot_date:
                continue
            item = rows_by_date.setdefault(snapshot_date, {
                "snapshot_date": snapshot_date,
                "metric_window": "daily",
            })
            item[metric] = values[1]
    return [rows_by_date[key] for key in sorted(rows_by_date)]


def read_xlsx_tables(path: Path) -> List[Tuple[str, List[Dict[str, Any]]]]:
    try:
        from openpyxl import load_workbook  # type: ignore
    except ImportError as exc:
        raise PipelineError("读取 Excel 需要 openpyxl；请先运行 python -m pip install -r requirements.txt") from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    tables: List[Tuple[str, List[Dict[str, Any]]]] = []
    sheet_rows: Dict[str, Sequence[Sequence[Any]]] = {}
    for sheet in workbook.worksheets:
        all_rows = list(sheet.iter_rows(values_only=True))
        sheet_rows[sheet.title] = all_rows
        header_index = _find_header_index(all_rows)
        if header_index is None:
            continue
        headers = all_rows[header_index]
        if not any(value not in (None, "") for value in headers):
            continue
        parsed = []
        for values in all_rows[header_index + 1:]:
            if not any(value not in (None, "") for value in values):
                continue
            parsed.append(canonicalize_row(dict(zip(headers, values))))
        if parsed:
            tables.append((sheet.title, parsed))
    daily_rows = _xhs_account_daily_rows(sheet_rows)
    if daily_rows:
        tables.append(("小红书账号每日趋势", daily_rows))
    return tables


def read_xls_tables(path: Path) -> List[Tuple[str, List[Dict[str, Any]]]]:
    try:
        import xlrd  # type: ignore
    except ImportError as exc:
        raise PipelineError("读取旧版 .xls 需要 xlrd；请先运行 python -m pip install -r requirements.txt") from exc
    workbook = xlrd.open_workbook(path)
    tables: List[Tuple[str, List[Dict[str, Any]]]] = []
    for sheet in workbook.sheets():
        if sheet.nrows < 2:
            continue
        all_rows = []
        for row_index in range(sheet.nrows):
            values = []
            for column_index in range(sheet.ncols):
                cell = sheet.cell(row_index, column_index)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    values.append(xlrd.xldate.xldate_as_datetime(cell.value, workbook.datemode))
                else:
                    values.append(cell.value)
            all_rows.append(values)
        header_index = _find_header_index(all_rows)
        if header_index is None:
            continue
        headers = all_rows[header_index]
        parsed = []
        for values in all_rows[header_index + 1:]:
            if not any(value not in (None, "") for value in values):
                continue
            parsed.append(canonicalize_row(dict(zip(headers, values))))
        if parsed:
            tables.append((sheet.name, parsed))
    return tables


def read_input_tables(path: Path) -> List[Tuple[str, List[Dict[str, Any]]]]:
    if path.suffix.lower() == ".csv":
        return [(path.stem, read_csv_rows(path))]
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        return read_xlsx_tables(path)
    if path.suffix.lower() == ".xls":
        return read_xls_tables(path)
    return []


def discover_imports(source: Path) -> List[Path]:
    if source.is_file():
        return [source]
    if not source.exists():
        raise PipelineError(f"导入路径不存在：{source}")
    files = [
        p for p in source.iterdir()
        if p.suffix.lower() in {".csv", ".xlsx", ".xlsm", ".xls"} and ".example." not in p.name
    ]
    if not files:
        raise PipelineError(f"未在 {source} 找到 CSV 或 Excel 文件。")
    return sorted(files)


def classify_table(name: str, rows: Sequence[Mapping[str, Any]]) -> str:
    keys = {key for row in rows[:5] for key in row}
    lowered = name.lower()
    note_identity = keys & {"content_id", "note_id", "title"}
    note_metrics = keys & set(METRICS)
    if note_identity and note_metrics:
        return "notes"
    if keys & {"followers", "followers_delta", "profile_views", "new_followers"}:
        return "account"
    if not note_identity and keys & {"snapshot_date", "period_start", "period_end"} and len(keys & set(ACCOUNT_METRICS)) >= 2:
        return "account"
    if ("note" in lowered or "笔记" in name or "作品" in name) and note_identity and note_metrics:
        return "notes"
    return "ignore"


def load_json(path: Path, default: Any = None) -> Any:
    if not path.exists():
        return default
    return json.loads(path.read_text(encoding="utf-8"))


def write_json_atomic(path: Path, payload: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    rendered = json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=False) + "\n"
    with tempfile.NamedTemporaryFile("w", encoding="utf-8", dir=path.parent, delete=False) as handle:
        handle.write(rendered)
        temporary = Path(handle.name)
    temporary.replace(path)


def load_content_master(project_root: Path) -> Dict[str, Any]:
    path = project_root / "data" / "content" / "content-master.json"
    master = load_json(path)
    if not master or "contents" not in master:
        raise PipelineError(f"内容主表不存在或格式错误：{path}")
    return master


def _content_indexes(master: Mapping[str, Any]) -> Tuple[Dict[str, Mapping[str, Any]], Dict[str, Mapping[str, Any]], Dict[str, Mapping[str, Any]]]:
    contents = master.get("contents", [])
    by_id = {str(item["content_id"]): item for item in contents}
    by_note = {str(item["xiaohongshu_note_id"]): item for item in contents if item.get("xiaohongshu_note_id")}
    by_title: Dict[str, Mapping[str, Any]] = {}
    for item in contents:
        for value in (item.get("title"), item.get("topic")):
            if value:
                by_title[str(value).strip()] = item
    return by_id, by_note, by_title


def normalize_note_rows(
    rows: Sequence[Mapping[str, Any]],
    master: Mapping[str, Any],
    default_date: str,
    unmatched_titles: Optional[List[str]] = None,
) -> List[Dict[str, Any]]:
    by_id, by_note, by_title = _content_indexes(master)
    normalized = []
    for index, raw in enumerate(rows, start=2):
        row = canonicalize_row(raw)
        content = None
        if row.get("content_id"):
            content = by_id.get(str(row["content_id"]).strip())
        if not content and row.get("note_id"):
            content = by_note.get(str(row["note_id"]).strip())
        if not content and row.get("title"):
            content = by_title.get(str(row["title"]).strip())
        if not content:
            if unmatched_titles is not None:
                unmatched_titles.append(str(row.get("title") or row.get("note_id") or row.get("content_id") or "未命名笔记"))
                continue
            raise PipelineError(f"笔记表第 {index} 行无法匹配内容主表：{row.get('title') or row.get('note_id') or row.get('content_id')}")
        snapshot_date = parse_date(row.get("snapshot_date"), "snapshot_date") or default_date
        publish_date = parse_date(row.get("publish_date"), "publish_date") or content.get("publish_date")
        item = {
            "snapshot_date": snapshot_date,
            "content_id": content["content_id"],
            "note_id": str(row.get("note_id") or content.get("xiaohongshu_note_id") or "") or None,
            "publish_date": publish_date,
        }
        for metric in METRICS:
            item[metric] = to_number(row.get(metric))
        normalized.append(item)
    return normalized


def normalize_account_rows(rows: Sequence[Mapping[str, Any]], default_date: str) -> List[Dict[str, Any]]:
    normalized = []
    for raw in rows:
        row = canonicalize_row(raw)
        period_start = parse_date(row.get("period_start"), "period_start")
        period_end = parse_date(row.get("period_end"), "period_end")
        raw_window = _header_key(row.get("metric_window"))
        window_aliases = {
            "daily": "daily", "单日": "daily", "每日": "daily", "日": "daily",
            "periodtotal": "period_total", "区间合计": "period_total", "周期合计": "period_total",
            "近7日": "period_total", "7日": "period_total", "30日": "period_total",
            "cumulative": "cumulative", "累计": "cumulative",
        }
        metric_window = window_aliases.get(raw_window)
        if metric_window is None:
            metric_window = "period_total" if period_start and period_end else "unknown"
        item = {
            "date": parse_date(row.get("snapshot_date"), "date") or period_end or default_date,
            "period_start": period_start,
            "period_end": period_end,
            "metric_window": metric_window,
        }
        for metric in ACCOUNT_METRICS:
            item[metric] = to_number(row.get(metric))
        normalized.append(item)
    return normalized


def _prior_snapshot(directory: Path, snapshot_date: str) -> Optional[Mapping[str, Any]]:
    candidates = sorted(p for p in directory.glob("*.json") if p.stem < snapshot_date)
    return load_json(candidates[-1]) if candidates else None


def validate_notes(notes: Sequence[Mapping[str, Any]], prior: Optional[Mapping[str, Any]]) -> Tuple[List[str], List[str]]:
    errors: List[str] = []
    warnings: List[str] = []
    if not notes:
        errors.append("笔记快照为空。")
        return errors, warnings
    ids = [item.get("content_id") for item in notes]
    if len(ids) != len(set(ids)):
        errors.append("同一天笔记快照含重复 content_id。")
    prior_notes = {item["content_id"]: item for item in (prior or {}).get("notes", [])}
    for metric in METRICS:
        present = [item.get(metric) for item in notes if item.get(metric) is not None]
        missing_count = sum(item.get(metric) is None for item in notes)
        if missing_count / len(notes) >= 0.5:
            warnings.append(f"{metric} 缺失 {missing_count}/{len(notes)}；缺失值保持 null。")
        if present and len(present) >= 3 and sum(value == 0 for value in present) / len(present) >= 0.7:
            previous_nonzero = any((item.get(metric) or 0) > 0 for item in prior_notes.values())
            message = f"{metric} 有 {sum(value == 0 for value in present)}/{len(present)} 个 0。"
            (errors if previous_nonzero else warnings).append(message)
    for item in notes:
        old = prior_notes.get(item["content_id"])
        if not old:
            continue
        for metric in METRICS:
            current, previous = item.get(metric), old.get(metric)
            if current is not None and previous is not None and current < previous:
                warnings.append(f"{item['content_id']} 的 {metric} 从 {previous} 降至 {current}，按平台修正处理，不计算负增量。")
    return errors, warnings


def validate_account(account: Mapping[str, Any], prior: Optional[Mapping[str, Any]]) -> Tuple[List[str], List[str]]:
    errors: List[str] = []
    warnings: List[str] = []
    if all(account.get(metric) is None for metric in ACCOUNT_METRICS):
        errors.append("账号快照没有可用指标。")
    missing_count = sum(account.get(metric) is None for metric in ACCOUNT_METRICS)
    if missing_count:
        warnings.append(f"账号快照缺失 {missing_count}/{len(ACCOUNT_METRICS)} 个可选指标；缺失值保持 null。")
    if account.get("metric_window") == "unknown":
        warnings.append("账号指标口径未声明；该快照不会用于近 7 日 KPI 汇总。")
    if account.get("metric_window") == "period_total":
        if not account.get("period_start") or not account.get("period_end"):
            errors.append("区间合计账号快照必须提供 period_start 和 period_end。")
        elif account["period_start"] > account["period_end"]:
            errors.append("账号快照 period_start 不得晚于 period_end。")
    old = (prior or {}).get("account", {})
    for metric in ("followers",):
        if account.get(metric) is not None and old.get(metric) is not None and account[metric] < old[metric]:
            warnings.append(f"账号 {metric} 从 {old[metric]} 降至 {account[metric]}，请确认是否为正常取关或平台修正。")
    return errors, warnings


def save_snapshot(path: Path, payload: Mapping[str, Any]) -> str:
    existing = load_json(path)
    if existing is not None:
        comparable_existing = dict(existing)
        comparable_new = dict(payload)
        for item in (comparable_existing, comparable_new):
            item.pop("imported_at", None)
            item.pop("source_files", None)
        if comparable_existing == comparable_new:
            return "unchanged"
        raise DuplicateSnapshotError(f"{path.name} 已存在且内容不同；为保护历史数据，本次未覆盖。")
    write_json_atomic(path, payload)
    return "created"


def _raw_root(project_root: Path) -> Path:
    publish_raw = os.getenv("PUBLISH_RAW_DATA", "false").lower() in {"1", "true", "yes"}
    return project_root / ("data/snapshots" if publish_raw else "local-data/raw")


def import_exports(project_root: Path, source: Path, import_date: Optional[str] = None) -> Dict[str, Any]:
    default_date = parse_date(import_date, "--date") if import_date else today_iso()
    assert default_date
    master = load_content_master(project_root)
    account_rows: List[Dict[str, Any]] = []
    note_rows: List[Dict[str, Any]] = []
    sources: List[str] = []
    warnings: List[str] = []
    unmatched_titles: List[str] = []
    for path in discover_imports(source):
        for table_name, rows in read_input_tables(path):
            if not rows:
                continue
            kind = classify_table(f"{path.name}:{table_name}", rows)
            if kind == "ignore":
                continue
            sources.append(f"{path.name}:{table_name}")
            if kind == "account":
                account_rows.extend(normalize_account_rows(rows, default_date))
            else:
                # The visible-table fallback has no stable note ID. Keep its raw
                # CSV intact, import matched rows, and surface new titles for the
                # content master instead of discarding the entire daily snapshot.
                allow_unmatched = path.name.startswith("visible-notes-table-")
                note_rows.extend(
                    normalize_note_rows(
                        rows,
                        master,
                        default_date,
                        unmatched_titles if allow_unmatched else None,
                    )
                )
    if unmatched_titles:
        warnings.append(
            f"页面采集有 {len(unmatched_titles)} 条笔记未匹配内容主表，已保留在私有 CSV、暂不进入分析："
            + "；".join(unmatched_titles)
        )
    if not account_rows and not note_rows:
        raise PipelineError("没有读取到可导入的数据。")

    raw_root = _raw_root(project_root)
    plan: List[Tuple[Path, Dict[str, Any]]] = []
    staged_prior: Optional[Mapping[str, Any]] = None
    for snapshot_date, grouped in sorted(_group_by(note_rows, "snapshot_date").items()):
        directory = raw_root / "notes"
        disk_prior = _prior_snapshot(directory, snapshot_date)
        prior = _latest_prior(disk_prior, staged_prior)
        errors, found_warnings = validate_notes(grouped, prior)
        warnings.extend(found_warnings)
        if errors:
            raise PipelineError("笔记数据校验失败：" + "；".join(errors))
        payload = {
            "schema_version": SCHEMA_VERSION,
            "snapshot_date": snapshot_date,
            "imported_at": now_iso(),
            "source_files": sources,
            "warnings": found_warnings,
            "notes": grouped,
        }
        plan.append((directory / f"{snapshot_date}.json", payload))
        staged_prior = payload
    staged_prior = None
    for snapshot_date, grouped in sorted(_group_by(account_rows, "date").items()):
        if len(grouped) != 1:
            raise PipelineError(f"账号快照 {snapshot_date} 有 {len(grouped)} 行，应为 1 行。")
        directory = raw_root / "account"
        disk_prior = _prior_snapshot(directory, snapshot_date)
        prior = _latest_prior(disk_prior, staged_prior)
        errors, found_warnings = validate_account(grouped[0], prior)
        warnings.extend(found_warnings)
        if errors:
            raise PipelineError("账号数据校验失败：" + "；".join(errors))
        payload = {
            "schema_version": SCHEMA_VERSION,
            "snapshot_date": snapshot_date,
            "imported_at": now_iso(),
            "source_files": sources,
            "warnings": found_warnings,
            "account": grouped[0],
        }
        plan.append((directory / f"{snapshot_date}.json", payload))
        staged_prior = payload

    # Validate the complete batch before making any write.
    for path, payload in plan:
        existing = load_json(path)
        if existing is not None:
            comparable_existing, comparable_new = dict(existing), dict(payload)
            for item in (comparable_existing, comparable_new):
                item.pop("imported_at", None)
                item.pop("source_files", None)
            if comparable_existing != comparable_new:
                raise DuplicateSnapshotError(f"{path.name} 已存在且内容不同；本批次未写入。")
    statuses = [{"path": str(path.relative_to(project_root)), "status": save_snapshot(path, payload)} for path, payload in plan]
    return {"date": default_date, "snapshots": statuses, "warnings": warnings, "raw_root": str(raw_root.relative_to(project_root))}


def _group_by(rows: Sequence[Mapping[str, Any]], key: str) -> Dict[str, List[Dict[str, Any]]]:
    result: Dict[str, List[Dict[str, Any]]] = defaultdict(list)
    for row in rows:
        result[str(row[key])].append(dict(row))
    return dict(result)


def _latest_prior(disk: Optional[Mapping[str, Any]], staged: Optional[Mapping[str, Any]]) -> Optional[Mapping[str, Any]]:
    candidates = [item for item in (disk, staged) if item]
    if not candidates:
        return None
    return max(candidates, key=lambda item: str(item.get("snapshot_date", "")))


def _load_snapshot_files(directory: Path, list_key: Optional[str] = None) -> List[Dict[str, Any]]:
    records: List[Dict[str, Any]] = []
    for path in sorted(directory.glob("*.json")):
        payload = load_json(path, {})
        if list_key:
            records.extend(payload.get(list_key, []))
        elif payload.get("account"):
            records.append(payload["account"])
    return records


def safe_delta(current: Optional[float], previous: Optional[float], label: str, warnings: List[str]) -> Optional[float]:
    if current is None or previous is None:
        return None
    delta = current - previous
    if delta < 0:
        warnings.append(f"{label} 累计值下降，增量记为 null。")
        return None
    return delta


def safe_rate(numerator: Optional[float], denominator: Optional[float]) -> Optional[float]:
    if numerator is None or denominator in (None, 0):
        return None
    return numerator / denominator


def metric_bundle(snapshot: Optional[Mapping[str, Any]]) -> Optional[Dict[str, Any]]:
    if snapshot is None:
        return None
    values = {metric: snapshot.get(metric) for metric in METRICS}
    interactions = None
    parts = [snapshot.get(metric) for metric in ("likes", "favorites", "comments", "shares")]
    if all(value is not None for value in parts):
        interactions = sum(parts)
    values.update({
        "interactions": interactions,
        "view_rate": safe_rate(snapshot.get("views"), snapshot.get("impressions")),
        "like_rate": safe_rate(snapshot.get("likes"), snapshot.get("views")),
        "favorite_rate": safe_rate(snapshot.get("favorites"), snapshot.get("views")),
        "comment_rate": safe_rate(snapshot.get("comments"), snapshot.get("views")),
        "share_rate": safe_rate(snapshot.get("shares"), snapshot.get("views")),
        "engagement_rate": safe_rate(interactions, snapshot.get("views")),
        "follow_conversion_rate": safe_rate(snapshot.get("followers_gained"), snapshot.get("views")),
        "snapshot_date": snapshot.get("snapshot_date"),
    })
    return values


def select_lifecycle_snapshot(snapshots: Sequence[Mapping[str, Any]], publish_date: Optional[str], day: int) -> Optional[Mapping[str, Any]]:
    if not publish_date:
        return None
    published = date.fromisoformat(publish_date)
    tolerance = 1 if day <= 3 else (2 if day == 7 else 3)
    eligible = []
    for snapshot in snapshots:
        age = (date.fromisoformat(str(snapshot["snapshot_date"])) - published).days
        if day <= age <= day + tolerance:
            eligible.append((age, snapshot))
    return min(eligible, key=lambda pair: pair[0])[1] if eligible else None


def _average(values: Iterable[Optional[float]]) -> Optional[float]:
    filtered = [float(value) for value in values if value is not None]
    return sum(filtered) / len(filtered) if filtered else None


def _median(values: Iterable[Optional[float]]) -> Optional[float]:
    filtered = [float(value) for value in values if value is not None]
    return statistics.median(filtered) if filtered else None


def _percentile_80(values: Sequence[float]) -> Optional[float]:
    if not values:
        return None
    ordered = sorted(values)
    return ordered[max(0, math.ceil(len(ordered) * 0.8) - 1)]


def _latest_by_content(note_rows: Sequence[Mapping[str, Any]]) -> Dict[str, Mapping[str, Any]]:
    latest: Dict[str, Mapping[str, Any]] = {}
    for row in sorted(note_rows, key=lambda item: str(item["snapshot_date"])):
        latest[str(row["content_id"])] = row
    return latest


def _rankings(content_results: Sequence[Mapping[str, Any]]) -> Dict[str, Any]:
    ranking_metrics = ("impressions", "views", "favorites", "interactions", "followers_gained")
    output: Dict[str, Any] = {}
    for window in (*WINDOWS.keys(), "cumulative"):
        output[window] = {}
        for metric in ranking_metrics:
            candidates = []
            for item in content_results:
                bundle = item["latest"] if window == "cumulative" else item["lifecycle"].get(window)
                value = bundle.get(metric) if bundle else None
                if value is not None:
                    candidates.append({"content_id": item["content_id"], "topic": item["topic"], "value": value})
            output[window][metric] = sorted(candidates, key=lambda item: item["value"], reverse=True)
    return output


def _series_rows(master: Mapping[str, Any], content_results: Sequence[Mapping[str, Any]]) -> List[Dict[str, Any]]:
    grouped: Dict[str, List[Mapping[str, Any]]] = defaultdict(list)
    for item in content_results:
        grouped[str(item["series"])].append(item)
    counts = {window: sum(item["lifecycle"].get(window) is not None for item in content_results) for window in WINDOWS}
    comparison_window = max((window for window in WINDOWS if counts[window]), key=lambda window: (counts[window], WINDOWS[window]), default=None)
    fair_views = []
    for item in content_results:
        bundle = item["lifecycle"].get(comparison_window) if comparison_window else None
        if bundle and bundle.get("views") is not None:
            fair_views.append(float(bundle["views"]))
    viral_threshold = _percentile_80(fair_views)
    output = []
    known_series = [item["name"] for item in master.get("series", [])]
    for series_name in [*known_series, *sorted(set(grouped) - set(known_series))]:
        items = grouped.get(series_name, [])
        fair_bundles = [item["lifecycle"].get(comparison_window) if comparison_window else None for item in items]
        fair_bundles = [bundle for bundle in fair_bundles if bundle]
        viral_count = sum(1 for bundle in fair_bundles if viral_threshold is not None and bundle.get("views") is not None and bundle["views"] >= viral_threshold)
        output.append({
            "series": series_name,
            "sample_size": len(fair_bundles),
            "average_impressions": _average(bundle.get("impressions") for bundle in fair_bundles),
            "median_impressions": _median(bundle.get("impressions") for bundle in fair_bundles),
            "average_views": _average(bundle.get("views") for bundle in fair_bundles),
            "average_favorite_rate": _average(bundle.get("favorite_rate") for bundle in fair_bundles),
            "average_engagement_rate": _average(bundle.get("engagement_rate") for bundle in fair_bundles),
            "average_followers_gained": _average(bundle.get("followers_gained") for bundle in fair_bundles),
            "follow_efficiency": _average(safe_rate(bundle.get("followers_gained"), bundle.get("views")) for bundle in fair_bundles),
            "long_tail_ability": _average(item.get("long_tail", {}).get("ratio") for item in items),
            "viral_rate": safe_rate(viral_count, len(fair_bundles)),
            "viral_threshold_views": viral_threshold,
            "comparison_window": comparison_window,
        })
    return output


def _topic_rows(master_by_id: Mapping[str, Mapping[str, Any]], content_results: Sequence[Mapping[str, Any]], dimension: str) -> List[Dict[str, Any]]:
    counts = {window: sum(item["lifecycle"].get(window) is not None for item in content_results) for window in WINDOWS}
    comparison_window = max((window for window in WINDOWS if counts[window]), key=lambda window: (counts[window], WINDOWS[window]), default=None)
    grouped: Dict[str, List[Mapping[str, Any]]] = defaultdict(list)
    for result in content_results:
        label = master_by_id[result["content_id"]].get(dimension) or "未分类"
        grouped[str(label)].append(result)
    output = []
    for label, items in sorted(grouped.items()):
        bundles = [item["lifecycle"].get(comparison_window) if comparison_window else None for item in items]
        bundles = [bundle for bundle in bundles if bundle]
        output.append({
            "value": label,
            "sample_size": len(bundles),
            "average_impressions": _average(bundle.get("impressions") for bundle in bundles),
            "average_views": _average(bundle.get("views") for bundle in bundles),
            "average_favorite_rate": _average(bundle.get("favorite_rate") for bundle in bundles),
            "average_engagement_rate": _average(bundle.get("engagement_rate") for bundle in bundles),
            "follow_efficiency": _average(safe_rate(bundle.get("followers_gained"), bundle.get("views")) for bundle in bundles),
            "comparison_window": comparison_window,
        })
    return output


def _account_summary(master: Mapping[str, Any], rows: Sequence[Mapping[str, Any]], fallback_date: Optional[str] = None) -> Dict[str, Any]:
    ordered = sorted(rows, key=lambda item: str(item["date"]))
    latest_date = str(ordered[-1]["date"]) if ordered else fallback_date
    latest = ordered[-1] if ordered else None
    if latest_date:
        end = date.fromisoformat(latest_date)
        start = end - timedelta(days=6)
        daily_rows = [item for item in ordered if item.get("metric_window") == "daily" and start <= date.fromisoformat(str(item["date"])) <= end]
        required_dates = {(start + timedelta(days=offset)).isoformat() for offset in range(7)}
        complete_dates = {str(item["date"]) for item in daily_rows} == required_dates
        period_rows = [
            item for item in ordered
            if item.get("metric_window") == "period_total"
            and item.get("period_start") == start.isoformat()
            and item.get("period_end") == end.isoformat()
        ]
        period_row = period_rows[-1] if period_rows else None
        published = [item for item in master.get("contents", []) if item.get("publish_date") and start <= date.fromisoformat(item["publish_date"]) <= end]

        def recent_value(metric: str) -> Optional[float]:
            if period_row is not None:
                return period_row.get(metric)
            if not complete_dates or any(item.get(metric) is None for item in daily_rows):
                return None
            return sum(item[metric] for item in daily_rows)

        if period_row is not None:
            follower_value = period_row.get("new_followers")
            if follower_value is None:
                follower_value = period_row.get("followers_delta")
            coverage_days = 7
            summary_window = "period_total"
        else:
            follower_values = [item.get("new_followers") if item.get("new_followers") is not None else item.get("followers_delta") for item in daily_rows]
            follower_value = sum(follower_values) if complete_dates and all(value is not None for value in follower_values) else None
            coverage_days = len({str(item["date"]) for item in daily_rows})
            summary_window = "daily" if complete_dates else None
        recent_summary = {
            "period_start": start.isoformat(),
            "period_end": end.isoformat(),
            "coverage_days": coverage_days,
            "metric_window": summary_window,
            "impressions": recent_value("impressions"),
            "views": recent_value("views"),
            "new_followers": follower_value,
            "publish_count": len(published),
        }
    else:
        recent_summary = {"period_start": None, "period_end": None, "coverage_days": 0, "metric_window": None, "impressions": None, "views": None, "new_followers": None, "publish_count": 0}
    return {
        "schema_version": SCHEMA_VERSION,
        "generated_at": now_iso(),
        "latest_date": latest_date,
        "latest": latest,
        "last_7_days": recent_summary,
        "trend": ordered,
    }


def _daily_insights(account_summary: Mapping[str, Any], series_rows: Sequence[Mapping[str, Any]], content_results: Sequence[Mapping[str, Any]], warnings: Sequence[str]) -> Dict[str, Any]:
    insights = []
    dated = account_summary.get("latest_date") or today_iso()
    eligible_series = [item for item in series_rows if item.get("sample_size", 0) >= 2 and item.get("average_favorite_rate") is not None]
    if eligible_series:
        best = max(eligible_series, key=lambda item: item["average_favorite_rate"])
        insights.append({"type": "positive", "title": f"{best['series']}收藏率领先", "detail": f"{best['sample_size']} 篇样本，平均收藏率 {best['average_favorite_rate']:.1%}。样本量仍需结合查看。"})
    long_tail = [item for item in content_results if (item.get("long_tail", {}).get("ratio") or 0) >= 0.3]
    if long_tail:
        top = max(long_tail, key=lambda item: item["long_tail"]["ratio"])
        insights.append({"type": "positive", "title": f"{top['topic']}出现持续长尾", "detail": f"发布 72 小时后的新增阅读占 7 日累计 {top['long_tail']['ratio']:.1%}。"})
    recent = account_summary.get("last_7_days", {})
    if recent.get("publish_count") == 0 and account_summary.get("latest_date"):
        insights.append({"type": "attention", "title": "近 7 日暂无发布", "detail": "发布频率为 0，请结合排期确认是否为主动停更。"})
    if warnings:
        insights.append({"type": "attention", "title": "发现数据修正或缺失", "detail": warnings[0]})
    if not insights:
        insights.append({"type": "neutral", "title": "等待更多快照", "detail": "导入至少两个采集日后，系统会识别增量、生命周期和异常变化。"})
    return {"schema_version": SCHEMA_VERSION, "generated_at": now_iso(), "date": dated, "insights": insights[:5]}


def analyze(project_root: Path) -> Dict[str, Any]:
    master = load_content_master(project_root)
    raw_root = _raw_root(project_root)
    note_rows = _load_snapshot_files(raw_root / "notes", "notes")
    account_rows = _load_snapshot_files(raw_root / "account")
    by_content: Dict[str, List[Mapping[str, Any]]] = defaultdict(list)
    for row in note_rows:
        by_content[str(row["content_id"])].append(row)
    master_by_id = {item["content_id"]: item for item in master.get("contents", [])}
    warnings: List[str] = []
    content_results = []
    for content_id, content in master_by_id.items():
        snapshots = sorted(by_content.get(content_id, []), key=lambda item: str(item["snapshot_date"]))
        latest = snapshots[-1] if snapshots else None
        previous = snapshots[-2] if len(snapshots) >= 2 else None
        lifecycle = {name: metric_bundle(select_lifecycle_snapshot(snapshots, content.get("publish_date"), days)) for name, days in WINDOWS.items()}
        bundle_latest = metric_bundle(latest)
        deltas = {}
        consecutive = bool(latest and previous and (date.fromisoformat(str(latest["snapshot_date"])) - date.fromisoformat(str(previous["snapshot_date"]))).days == 1)
        if latest and previous and not consecutive:
            warnings.append(f"{content_id} 最近两个快照不连续，不计算日增量。")
        for metric in METRICS:
            deltas[metric] = safe_delta(latest.get(metric), previous.get(metric), f"{content_id} {metric}", warnings) if consecutive else None
        views_72 = lifecycle["72h"].get("views") if lifecycle["72h"] else None
        views_7d = lifecycle["7d"].get("views") if lifecycle["7d"] else None
        long_tail_ratio = safe_rate((views_7d - views_72) if views_7d is not None and views_72 is not None and views_7d >= views_72 else None, views_7d)
        content_results.append({
            "content_id": content_id,
            "title": content.get("title"),
            "topic": content.get("topic"),
            "series": content.get("series"),
            "publish_date": content.get("publish_date"),
            "latest": bundle_latest,
            "deltas": deltas,
            "lifecycle": lifecycle,
            "long_tail": {"metric": "7d_after_72h_ratio", "ratio": long_tail_ratio},
        })
    latest_note_date = max((str(row["snapshot_date"]) for row in note_rows), default=None)
    account_summary = _account_summary(master, account_rows, latest_note_date)
    series_rows = _series_rows(master, content_results)
    topic_dimensions = {dimension: _topic_rows(master_by_id, content_results, dimension) for dimension in ("content_type", "title_type", "industry", "series")}
    generated_at = now_iso()
    derived = project_root / "data" / "derived"
    content_payload = {"schema_version": SCHEMA_VERSION, "generated_at": generated_at, "warnings": warnings, "contents": content_results, "rankings": _rankings(content_results)}
    series_payload = {"schema_version": SCHEMA_VERSION, "generated_at": generated_at, "series": series_rows}
    topic_payload = {"schema_version": SCHEMA_VERSION, "generated_at": generated_at, "dimensions": topic_dimensions}
    insights_payload = _daily_insights(account_summary, series_rows, content_results, warnings)
    outputs = {
        "account-summary.json": account_summary,
        "content-metrics.json": content_payload,
        "series-metrics.json": series_payload,
        "topic-metrics.json": topic_payload,
        "daily-insights.json": insights_payload,
    }
    for filename, payload in outputs.items():
        write_json_atomic(derived / filename, payload)
    return {"generated_at": generated_at, "note_snapshots": len(note_rows), "account_snapshots": len(account_rows), "warnings": warnings, "outputs": [str((derived / name).relative_to(project_root)) for name in outputs]}
